import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# Enter the file path to your raw_financials.xlsx file here
# Example: "../folder_name/raw_financials.xlsx"
RAW_FILE_PATH = "ENTER YOUR FILE PATH HERE"

# Enter the file path where you want to save the completed financial model
# Example: "../folder_name/financial_model.xlsx"
OUTPUT_FILE_PATH = "ENTER YOUR OUTPUT FILE PATH HERE"

# ============================================================
# A. RAW MODEL FINANCIALS
# ============================================================

df_A = pd.read_excel(
    RAW_FILE_PATH,
    sheet_name="Model Financials"
)
# ============================================================
# B. HISTORICAL MODEL FINANCIALS
# ============================================================

# B is a separate copy of A.
# We keep all original historical data and add our
# calculated historical metrics to this dataset.

df_B = df_A.copy()

# =========================
# HISTORICAL CALCULATIONS
# =========================


# -------------------------
# Income Statement
# -------------------------

# EBIT = Revenue - COGS - SG&A - D&A

df_B["EBIT"] = (
    df_B["Revenue"]
    - df_B["COGS"]
    - df_B["SG&A"]
    - df_B["D&A"]
)


# EBT = EBIT - Interest Expense

df_B["EBT"] = (
    df_B["EBIT"]
    - df_B["Interest Expense"]
)


# Tax Rate

df_B["Tax Rate"] = (
    df_B["Tax Expense"]
    / df_B["EBT"]
    * 100
)


# -------------------------
# Cash Flow Statement
# -------------------------

# Changes in working capital

change_in_AR = (
    df_B["Accounts Receivable"]
    .diff()
    .fillna(0)
)

change_in_inventory = (
    df_B["Inventory"]
    .diff()
    .fillna(0)
)

change_in_AP = (
    df_B["Accounts Payable"]
    .diff()
    .fillna(0)
)


# Operating Cash Flow

df_B["Operating Cash Flow"] = (
    df_B["Net Income"]
    + df_B["D&A"]
    - change_in_AR
    - change_in_inventory
    + change_in_AP
)


# Investing Cash Flow
# In our model, CapEx is the only investing activity.

df_B["Investing Cash Flow"] = (
    -df_B["Capital Expenditure"]
)


# Changes in debt

change_in_short_term_debt = (
    df_B["Short-Term Debt"]
    .diff()
    .fillna(0)
)

change_in_long_term_debt = (
    df_B["Long-Term Debt"]
    .diff()
    .fillna(0)
)


# Debt financing

debt_financing = (
    change_in_short_term_debt
    + change_in_long_term_debt
)


# Equity financing

equity_financing = (
    df_B["New Equity Issued"]
    - df_B["Dividends Paid"]
)


# Financing Cash Flow

df_B["Financing Cash Flow"] = (
    debt_financing
    + equity_financing
)


# Change in Cash

df_B["Change in Cash"] = (
    df_B["Operating Cash Flow"]
    + df_B["Investing Cash Flow"]
    + df_B["Financing Cash Flow"]
)
# Create Cash column
df_B["Cash"] = 0.0

# Starting cash balance for 2021
df_B.loc[0, "Cash"] = 1100

# Calculate ending Cash for 2022–2025
for i in range(1, len(df_B)):
    df_B.loc[i, "Cash"] = (
        df_B.loc[i - 1, "Cash"]
        + df_B.loc[i, "Change in Cash"]
    )

# -------------------------
# Balance Sheet
# -------------------------

# Total Assets

df_B["Total Assets"] = (
    df_B["Cash"]
    + df_B["Accounts Receivable"]
    + df_B["Inventory"]
    + df_B["Ending PP&E"]
)


# Total Debt

df_B["Total Debt"] = (
    df_B["Short-Term Debt"]
    + df_B["Long-Term Debt"]
)


# Interest Rate

df_B["Interest Rate"] = (
    df_B["Interest Expense"]
    / df_B["Total Debt"]
    * 100
)


# Total Liabilities

df_B["Total Liabilities"] = (
    df_B["Accounts Payable"]
    + df_B["Short-Term Debt"]
    + df_B["Long-Term Debt"]
)


# Total Liabilities & Equity

df_B["Total Liabilities & Equity"] = (
    df_B["Total Liabilities"]
    + df_B["Ending Equity"]
)


# Balance Sheet Check

df_B["Balance Check"] = (
    df_B["Total Assets"]
    - df_B["Total Liabilities & Equity"]
)


# ============================================================
# FORMAT HISTORICAL SHEET
# ============================================================


df_B.to_excel(
    OUTPUT_FILE_PATH,
    sheet_name="Historical",
    index=False
)
wb = load_workbook(OUTPUT_FILE_PATH)

ws = wb["Historical"]

ws.freeze_panes = "B2"



for cell in ws[1]:

    cell.font = Font(bold=True, color="FFFFFF")

    cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")

    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


for row in ws.iter_rows(min_row=2, min_col=2):
    for cell in row:
        cell.number_format = "0.00"


for column in ws.columns:

    max_length = 0



    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column_letter].width = max_length + 4


# ============================================================
# SAVE HISTORICAL MODEL
# ============================================================

wb.save(OUTPUT_FILE_PATH)

# ============================================================
# C. FORECAST MODEL FINANCIALS
# ============================================================

model_columns = df_B.columns.tolist()

# ============================================================
# FORECAST ASSUMPTIONS
# ============================================================

# -------------------------
# Revenue Growth
# -------------------------

historical_avg_yoy = (
    df_B["Revenue"].pct_change().mean() * 100
)

recent_CAGR = (
    (df_B["Revenue"].iloc[-1] / df_B["Revenue"].iloc[-4])
    ** (1 / 3)
    - 1
) * 100

recent_momentum = (
    df_B["Revenue"].pct_change().iloc[-1] * 100
)

final_revenue_growth = (
    recent_momentum * 0.50
    + recent_CAGR * 0.30
    + historical_avg_yoy * 0.20
)


# -------------------------
# COGS
# -------------------------

COGS_REV = (
    df_B["COGS"]
    / df_B["Revenue"]
    * 100
)

average_COGS_REV = COGS_REV.mean()
recent_COGS_REV = COGS_REV.iloc[-1]

final_COGS_REV = (
    average_COGS_REV * 0.50
    + recent_COGS_REV * 0.50
)


# -------------------------
# SG&A
# -------------------------

SGA_REV = (
    df_B["SG&A"]
    / df_B["Revenue"]
    * 100
)

average_SGA_REV = SGA_REV.mean()
recent_SGA_REV = SGA_REV.iloc[-1]

final_SGA_REV = (
    average_SGA_REV * 0.50
    + recent_SGA_REV * 0.50
)


# -------------------------
# D&A
# -------------------------

DA_PPE = (
    df_B["D&A"]
    / df_B["Beginning PP&E"]
    * 100
)

average_DA_PPE = DA_PPE.mean()
recent_DA_PPE = DA_PPE.iloc[-1]

final_DA_PPE = (
    average_DA_PPE * 0.50
    + recent_DA_PPE * 0.50
)


# -------------------------
# CapEx
# -------------------------

CapEx_PPE = (
    df_B["Capital Expenditure"]
    / df_B["Beginning PP&E"]
    * 100
)

average_CapEx_PPE = CapEx_PPE.mean()
recent_CapEx_PPE = CapEx_PPE.iloc[-1]

final_CapEx_PPE = (
    average_CapEx_PPE * 0.50
    + recent_CapEx_PPE * 0.50
)


# -------------------------
# Interest Rate
# -------------------------

interest_rate_average = (
    df_B["Interest Rate"].mean()
)

recent_interest_rate = (
    df_B["Interest Rate"].iloc[-1]
)

final_interest_rate = (
    interest_rate_average * 0.50
    + recent_interest_rate * 0.50
)


# -------------------------
# Debt Growth
# -------------------------

Debt_Growth = (
    df_B["Total Debt"].pct_change()
    * 100
)

recent_debt_growth = (
    Debt_Growth.iloc[-1]
)

average_debt_growth = (
    Debt_Growth.mean()
)

recent_debt_CAGR = (
    (df_B["Total Debt"].iloc[-1] / df_B["Total Debt"].iloc[-4])
    ** (1 / 3)
    - 1
) * 100

final_debt_growth = (
    recent_debt_growth * 0.50
    + recent_debt_CAGR * 0.30
    + average_debt_growth * 0.20
)


# -------------------------
# Tax Rate
# -------------------------

tax_rate_average = (
    df_B["Tax Rate"].mean()
)

recent_tax_rate = (
    df_B["Tax Rate"].iloc[-1]
)

final_tax_rate = (
    tax_rate_average * 0.50
    + recent_tax_rate * 0.50
)

# -------------------------
# Accounts Receivable
# -------------------------

AR_REV = (
    df_B["Accounts Receivable"]
    / df_B["Revenue"]
    * 100
)

average_AR_REV = AR_REV.mean()
recent_AR_REV = AR_REV.iloc[-1]

final_AR_REV = (
    average_AR_REV * 0.50
    + recent_AR_REV * 0.50
)


# -------------------------
# Inventory
# -------------------------

INV_REV = (
    df_B["Inventory"]
    / df_B["Revenue"]
    * 100
)

average_INV_REV = INV_REV.mean()
recent_INV_REV = INV_REV.iloc[-1]

final_Inventory_REV = (
    average_INV_REV * 0.50
    + recent_INV_REV * 0.50
)


# -------------------------
# Accounts Payable
# -------------------------

AP_REV = (
    df_B["Accounts Payable"]
    / df_B["Revenue"]
    * 100
)

average_AP_REV = AP_REV.mean()
recent_AP_REV = AP_REV.iloc[-1]

final_AP_REV = (
    average_AP_REV * 0.50
    + recent_AP_REV * 0.50
)


# -------------------------
# Short-Term Debt
# -------------------------

STD_TD = (
    df_B["Short-Term Debt"]
    / df_B["Total Debt"]
    * 100
)

average_STD_DEBT = STD_TD.mean()
recent_STD_DEBT = STD_TD.iloc[-1]

final_STD_DEBT = (
    average_STD_DEBT * 0.50
    + recent_STD_DEBT * 0.50
)


# ============================================================
# FORECAST LOOP
# ============================================================

# Starting values come from the final historical year, 2025.

previous_revenue = (
    df_B["Revenue"].iloc[-1]
)

beginning_PPE = (
    df_B["Ending PP&E"].iloc[-1]
)

previous_debt = (
    df_B["Total Debt"].iloc[-1]
)

beginning_equity = (
    df_B["Ending Equity"].iloc[-1]
)

previous_cash = (
    df_B["Cash"].iloc[-1]
)

previous_AR = (
    df_B["Accounts Receivable"].iloc[-1]
)

previous_Inventory = (
    df_B["Inventory"].iloc[-1]
)

previous_AP = (
    df_B["Accounts Payable"].iloc[-1]
)


# Forecast years

forecast_years = [
    2026,
    2027,
    2028,
    2029,
    2030
]


# Empty list to store one dictionary for each forecast year

forecast_results = []


# ------------------------------------------------------------
# Start Forecast Loop
# ------------------------------------------------------------

for forecast_year in forecast_years:

    # -------------------------
    # Revenue
    # -------------------------

    forecasted_revenue = (
        previous_revenue
        * (1 + final_revenue_growth / 100)
    )


    # -------------------------
    # COGS
    # -------------------------

    forecasted_COGS = (
        final_COGS_REV
        / 100
        * forecasted_revenue
    )


    # -------------------------
    # SG&A
    # -------------------------

    forecasted_SGA = (
        final_SGA_REV
        / 100
        * forecasted_revenue
    )


    # -------------------------
    # D&A
    # -------------------------

    forecasted_DA = (
        beginning_PPE
        * final_DA_PPE
        / 100
    )


    # -------------------------
    # Capital Expenditure
    # -------------------------

    forecasted_CapEx = (
        beginning_PPE
        * final_CapEx_PPE
        / 100
    )


    # -------------------------
    # Ending PP&E
    # -------------------------

    forecasted_ending_PPE = (
        beginning_PPE
        + forecasted_CapEx
        - forecasted_DA
    )


    # -------------------------
    # Total Debt
    # -------------------------

    forecasted_debt = (
        previous_debt
        * (1 + final_debt_growth / 100)
    )


    # -------------------------
    # Interest Expense
    # -------------------------

    forecasted_interest_expense = (
        forecasted_debt
        * final_interest_rate
        / 100
    )


    # -------------------------
    # EBIT
    # -------------------------

    forecasted_EBIT = (
        forecasted_revenue
        - forecasted_COGS
        - forecasted_SGA
        - forecasted_DA
    )


    # -------------------------
    # EBT
    # -------------------------

    forecasted_EBT = (
        forecasted_EBIT
        - forecasted_interest_expense
    )


    # -------------------------
    # Tax Expense
    # -------------------------

    forecasted_tax_expense = (
        forecasted_EBT
        * final_tax_rate
        / 100
    )


    # -------------------------
    # Net Income
    # -------------------------

    forecasted_net_income = (
        forecasted_EBT
        - forecasted_tax_expense
    )


    # -------------------------
    # Accounts Receivable
    # -------------------------

    forecasted_AR = (
        forecasted_revenue
        * final_AR_REV
        / 100
    )


    # -------------------------
    # Inventory
    # -------------------------

    forecasted_Inventory = (
        forecasted_revenue
        * final_Inventory_REV
        / 100
    )


    # -------------------------
    # Accounts Payable
    # -------------------------

    forecasted_AP = (
        forecasted_revenue
        * final_AP_REV
        / 100
    )


    # -------------------------
    # Short-Term Debt
    # -------------------------

    forecasted_STD = (
        forecasted_debt
        * final_STD_DEBT
        / 100
    )


    # -------------------------
    # Long-Term Debt
    # -------------------------

    forecasted_LTD = (
        forecasted_debt
        - forecasted_STD
    )


    # -------------------------
    # Equity
    # -------------------------

    forecasted_dividends = (
        forecasted_net_income
        * 0.80
    )

    forecasted_new_equity = 0

    forecasted_ending_equity = (
        beginning_equity
        + forecasted_net_income
        - forecasted_dividends
        + forecasted_new_equity
    )


    # -------------------------
    # Working Capital Changes
    # -------------------------

    change_in_AR = (
        forecasted_AR
        - previous_AR
    )

    change_in_inventory = (
        forecasted_Inventory
        - previous_Inventory
    )

    change_in_AP = (
        forecasted_AP
        - previous_AP
    )


    # -------------------------
    # Operating Cash Flow
    # -------------------------

    forecasted_operating_cash_flow = (
        forecasted_net_income
        + forecasted_DA
        - change_in_AR
        - change_in_inventory
        + change_in_AP
    )


    # -------------------------
    # Investing Cash Flow
    # -------------------------

    forecasted_investing_cash_flow = (
        -forecasted_CapEx
    )


    # -------------------------
    # Financing Cash Flow
    # -------------------------

    forecasted_debt_financing = (
        forecasted_debt
        - previous_debt
    )

    forecasted_equity_financing = (
        forecasted_new_equity
        - forecasted_dividends
    )

    forecasted_financing_cash_flow = (
        forecasted_debt_financing
        + forecasted_equity_financing
    )


    # -------------------------
    # Change in Cash
    # -------------------------

    forecasted_change_in_cash = (
        forecasted_operating_cash_flow
        + forecasted_investing_cash_flow
        + forecasted_financing_cash_flow
    )


    # -------------------------
    # Ending Cash
    # -------------------------

    forecasted_cash = (
        previous_cash
        + forecasted_change_in_cash
    )


    # -------------------------
    # Total Assets
    # -------------------------

    forecasted_total_assets = (
        forecasted_cash
        + forecasted_AR
        + forecasted_Inventory
        + forecasted_ending_PPE
    )


    # -------------------------
    # Total Liabilities
    # -------------------------

    forecasted_total_liabilities = (
        forecasted_AP
        + forecasted_STD
        + forecasted_LTD
    )


    # -------------------------
    # Total Liabilities & Equity
    # -------------------------

    forecasted_total_liabilities_equity = (
        forecasted_total_liabilities
        + forecasted_ending_equity
    )


    # -------------------------
    # Interest Rate
    # -------------------------

    forecasted_interest_rate = (
        forecasted_interest_expense
        / forecasted_debt
        * 100
    )


    # -------------------------
    # Tax Rate
    # -------------------------

    forecasted_tax_rate = (
        forecasted_tax_expense
        / forecasted_EBT
        * 100
    )


    # -------------------------
    # PP&E Check
    # -------------------------

    forecasted_PPE_check = (
        beginning_PPE
        + forecasted_CapEx
        - forecasted_DA
        - forecasted_ending_PPE
    )


    # -------------------------
    # Equity Check
    # -------------------------

    forecasted_equity_check = (
        beginning_equity
        + forecasted_net_income
        - forecasted_dividends
        + forecasted_new_equity
        - forecasted_ending_equity
    )


    # -------------------------
    # Balance Sheet Check
    # -------------------------

    forecasted_balance_check = round(
        forecasted_total_assets
        - forecasted_total_liabilities_equity, 10)


    # ========================================================
    # STORE FORECAST YEAR
    # ========================================================

    forecast_results.append({

        # Income Statement
        "Year": forecast_year,
        "Revenue": forecasted_revenue,
        "COGS": forecasted_COGS,
        "SG&A": forecasted_SGA,
        "D&A": forecasted_DA,
        "EBIT": forecasted_EBIT,
        "Interest Expense": forecasted_interest_expense,
        "EBT": forecasted_EBT,
        "Tax Rate": forecasted_tax_rate,
        "Tax Expense": forecasted_tax_expense,
        "Net Income": forecasted_net_income,

        # Balance Sheet
        "Cash": forecasted_cash,
        "Accounts Receivable": forecasted_AR,
        "Inventory": forecasted_Inventory,
        "Beginning PP&E": beginning_PPE,
        "Capital Expenditure": forecasted_CapEx,
        "Ending PP&E": forecasted_ending_PPE,
        "Accounts Payable": forecasted_AP,

        "Short-Term Debt": forecasted_STD,
        "Long-Term Debt": forecasted_LTD,
        "Total Debt": forecasted_debt,
        "Interest Rate": forecasted_interest_rate,

        "Beginning Equity": beginning_equity,
        "Ending Equity": forecasted_ending_equity,
        "Dividends Paid": forecasted_dividends,
        "New Equity Issued": forecasted_new_equity,

        # Cash Flow Statement
        "Operating Cash Flow": forecasted_operating_cash_flow,
        "Investing Cash Flow": forecasted_investing_cash_flow,
        "Financing Cash Flow": forecasted_financing_cash_flow,
        "Change in Cash": forecasted_change_in_cash,

        # Balance Sheet Totals
        "Total Assets": forecasted_total_assets,
        "Total Liabilities": forecasted_total_liabilities,
        "Total Liabilities & Equity": forecasted_total_liabilities_equity,

        # Checks
        "PP&E Check": forecasted_PPE_check,
        "Equity Check": forecasted_equity_check,
        "Balance Check": forecasted_balance_check
    })


    # ========================================================
    # UPDATE ROLL-FORWARD VALUES FOR NEXT YEAR
    # ========================================================

    previous_revenue = forecasted_revenue

    beginning_PPE = forecasted_ending_PPE

    previous_debt = forecasted_debt

    beginning_equity = forecasted_ending_equity

    previous_cash = forecasted_cash

    previous_AR = forecasted_AR

    previous_Inventory = forecasted_Inventory

    previous_AP = forecasted_AP


# ============================================================
# CREATE FORECAST DATASET
# ============================================================

df_C = pd.DataFrame(forecast_results)



df_C = df_C[model_columns]

# ============================================================
# SAVE FORECAST TO EXCEL
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE_PATH,
    engine="openpyxl",
    mode="a"
) as writer:

    df_C.to_excel(
        writer,
        sheet_name="Forecast",
        index=False
    )


# ============================================================
# FORMAT FORECAST SHEET
# ============================================================

wb = load_workbook(OUTPUT_FILE_PATH)

ws = wb["Forecast"]

ws.freeze_panes = "B2"
for cell in ws[1]:
    cell.font = Font(bold=True,color="FFFFFF")
    cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


for row in ws.iter_rows(min_row=2, min_col=2):
    for cell in row:
        cell.number_format = "0.00"



for column in ws.columns:

    max_length = 0



    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column_letter].width = max_length + 4


# ============================================================
# SAVE FORECAST MODEL
# ============================================================


wb.save(OUTPUT_FILE_PATH)

# ============================================================
# DCF
# ============================================================

# -------------------------
# NWC/Change in NWC
# -------------------------

NWC_2025 = (
    df_B["Accounts Receivable"].iloc[-1]
    + df_B["Inventory"].iloc[-1]
    - df_B["Accounts Payable"].iloc[-1]
)

df_D = df_C.copy()

df_D["NWC"] = (
    df_D["Accounts Receivable"]
    + df_D["Inventory"]
    - df_D["Accounts Payable"]
)

change_NWC_2026 = df_D.loc[0, "NWC"] - NWC_2025

df_D["Change in NWC"] = df_D["NWC"].diff().fillna(change_NWC_2026)


# -------------------------
# NOPAT (Net Operating Profit After Tax)
# -------------------------

df_D["Taxes on EBIT"] = (
    df_D["EBIT"]
    * df_D["Tax Rate"]
    / 100
)

df_D["NOPAT"] = (
    df_D["EBIT"]
    - df_D["Taxes on EBIT"]
)

# -------------------------
# UFCF (Unlevered Free Cash Flow)
# -------------------------

df_D["UFCF"] = (
    df_D["NOPAT"]
    + df_D["D&A"]
    - df_D["Capital Expenditure"]
    - df_D["Change in NWC"]
)
# -------------------------
# WACC — Weighted Average Cost of Capital
# -------------------------

# WACC =
# (Market value of Equity / Total Capital) × Cost of Equity
# +
# (Market value of Debt / Total Capital) × Cost of Debt × (1 − Tax Rate)

# /////////////////////////////////////

# DCF Assumptions

# Market value of equity is assumed because the company is fictional
# and therefore has no observable market capitalization.
market_value_of_equity = 8000

# Market value of debt is assumed to equal 2025 Total Debt
# because the company is fictional and has no observable debt market value.
market_value_of_debt = df_B["Total Debt"].iloc[-1]

# Assumed risk-free rate.
risk_free_rate = 6.5

# Assumed beta for the fictional company.
beta = 1.1

# Assumed equity risk premium.
equity_risk_premium = 5.5

# /////////////////////////////////////

# WACC Capital Structure

# Total capital is the combined market value of equity and debt.
total_capital = market_value_of_equity + market_value_of_debt

# These weights represent the proportion of total capital
# financed by equity and debt.
equity_weight = market_value_of_equity / total_capital
debt_weight = market_value_of_debt / total_capital

# /////////////////////////////////////

# Cost of Debt

# The pre-tax cost of debt is based on the forecast interest rate
# calculated from the company's historical interest-rate methodology.
cost_of_debt = final_interest_rate


# Cost of Equity — CAPM


# CAPM is used to estimate the return required by equity investors.
cost_of_equity = (
    risk_free_rate
    + beta * equity_risk_premium
)

# After-tax cost of debt

after_tax_cost_of_debt = (
    cost_of_debt
    * (1 - final_tax_rate / 100)
)

# /////////////////////////////////////

# WACC

wacc = (
    equity_weight * cost_of_equity
    + debt_weight * after_tax_cost_of_debt
)

# -------------------------
# Discounting each year's UFCF
# -------------------------

# PV(UFCFt) = UFCFt/(1+WACC)^t

df_D["Discount Period"] = range(1, len(df_D) + 1)

df_D["PV of UFCF"] = (
    df_D["UFCF"]
    / (1 + wacc / 100) ** df_D["Discount Period"]
)

# -------------------------
# Terminal Value (Gordon Growth Model)
# -------------------------

# UFCF Growth (Long-term perpetual growth rate)

terminal_growth_rate = 3.0

# Terminal Value

terminal_value = (
    df_D["UFCF"].iloc[-1]
    * (1 + terminal_growth_rate / 100)
    / ((wacc / 100) - (terminal_growth_rate / 100))
)

# Present Value of Terminal Value

pv_terminal_value = (
    terminal_value
    / (1 + wacc / 100) ** 5
)

# /////////////////////////////////////

# Present Value of Explicit Forecast

pv_sum_of_UFCF = df_D["PV of UFCF"].sum()
# Enterprise Value

enterprise_value = (
    pv_sum_of_UFCF
    + pv_terminal_value
)

# Equity Value

equity_value = (
    enterprise_value
    - market_value_of_debt
    + df_B["Cash"].iloc[-1]
)
# /////////////////////////////////////

# -------------------------
# Implied Share Price
# -------------------------

# Assumed shares outstanding for the fictional company.
shares_outstanding = 1000

# Implied Share Price

implied_share_price = (
    equity_value / shares_outstanding
)

# -------------------------
# Sensitivity Analysis
# -------------------------

wacc_sensitivity = [wacc - 1, wacc, wacc + 1]
terminal_growth_sensitivity = [2, 3, 4]


sensitivity_results = []


for sensitivity_wacc in wacc_sensitivity:
    for sensitivity_growth in terminal_growth_sensitivity:
        pv_sum = (
            df_D["UFCF"]
            / (1 + sensitivity_wacc / 100) ** df_D["Discount Period"]
        ).sum()

        sensitivity_terminal_value = (
            df_D["UFCF"].iloc[-1]
            * (1 + sensitivity_growth / 100)
            / (
                (sensitivity_wacc / 100)
                - (sensitivity_growth / 100)
            )
        )

        sensitivity_pv_terminal_value = (
            sensitivity_terminal_value
            / (1 + sensitivity_wacc / 100) ** 5
        )

        sensitivity_enterprise_value = (
            pv_sum
            + sensitivity_pv_terminal_value
        )

        sensitivity_equity_value = (
            sensitivity_enterprise_value
            - market_value_of_debt
            + df_B["Cash"].iloc[-1]
        )

        sensitivity_share_price = (
            sensitivity_equity_value
            / shares_outstanding
        )
        sensitivity_results.append({
            "WACC": sensitivity_wacc,
            "Terminal Growth": sensitivity_growth,
            "Implied Share Price": sensitivity_share_price
        })

sensitivity_df = pd.DataFrame(sensitivity_results)

sensitivity_table = sensitivity_df.pivot(
    index="WACC",
    columns="Terminal Growth",
    values="Implied Share Price"
)

# ============================================================
# SAVE DCF TO EXCEL
# ============================================================

dcf_columns = [
    "Year",
    "NWC",
    "Change in NWC",
    "Taxes on EBIT",
    "NOPAT",
    "UFCF",
    "Discount Period",
    "PV of UFCF"
]

summary_data = [
    ["DCF Assumptions", ""],
    ["Market Value of Equity", market_value_of_equity],
    ["Market Value of Debt", market_value_of_debt],
    ["Risk-Free Rate", risk_free_rate],
    ["Beta", beta],
    ["Equity Risk Premium", equity_risk_premium],
    ["Terminal Growth Rate", terminal_growth_rate],
    ["Shares Outstanding", shares_outstanding],
    ["", ""],
    ["DCF Valuation", ""],
    ["Cost of Equity", cost_of_equity],
    ["Cost of Debt", cost_of_debt],
    ["After-Tax Cost of Debt", after_tax_cost_of_debt],
    ["WACC", wacc],
    ["Sum of PV of Forecast UFCF", pv_sum_of_UFCF],
    ["PV of Terminal Value", pv_terminal_value],
    ["Enterprise Value", enterprise_value],
    ["Equity Value", equity_value],
    ["Implied Share Price", implied_share_price]
]

summary_df = pd.DataFrame(
    summary_data,
    columns=["Metric", "Value"]
)


# Export DCF, Sensitivity Analysis, and Summary

# Overlay allows the sensitivity table and DCF summary
# to be written into the same DCF worksheet.

with pd.ExcelWriter(
    OUTPUT_FILE_PATH,
    engine="openpyxl",
    mode="a",
    if_sheet_exists="overlay"
) as writer:

    df_D[dcf_columns].to_excel(
        writer,
        sheet_name="DCF",
        index=False
    )

    sensitivity_table.to_excel(
        writer,
        sheet_name="DCF",
        startrow=len(df_D) + 4
    )

    summary_df.to_excel(
        writer,
        sheet_name="DCF",
        index=False,
        startrow=len(df_D) + 9
    )


# Format DCF sheet
wb = load_workbook(OUTPUT_FILE_PATH)
ws = wb["DCF"]

ws.freeze_panes = "B2"


# Format DCF table header
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# Center-align DCF table data
for row in ws.iter_rows(
    min_row=2,
    max_row=len(df_D) + 1
):
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# Format DCF table numbers
for row in ws.iter_rows(
    min_row=2,
    max_row=len(df_D) + 1,
    min_col=2
):
    for cell in row:
        cell.number_format = "0.00"


# Sensitivity Analysis title

sensitivity_title_row = len(df_D) + 4

ws.cell(
    row=sensitivity_title_row,
    column=1,
    value="Sensitivity Analysis"
)

ws.cell(
    row=sensitivity_title_row,
    column=1
).font = Font(bold=True)

# Terminal Growth Rate header
ws.merge_cells(
    start_row=sensitivity_title_row,
    start_column=2,
    end_row=sensitivity_title_row,
    end_column=4
)

ws.cell(
    row=sensitivity_title_row,
    column=2,
    value="Terminal Growth Rate (%)"
)

ws.cell(
    row=sensitivity_title_row,
    column=2
).font = Font(
    bold=True,
    color="FFFFFF"
)

ws.cell(
    row=sensitivity_title_row,
    column=2
).alignment = Alignment(
    horizontal="center",
    vertical="center"
)

# Blue background across merged B:D header
for col in range(2, 5):
    ws.cell(
        row=sensitivity_title_row,
        column=col
    ).fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

# Sensitivity Analysis header
sensitivity_header_row = len(df_D) + 5

for cell in ws[sensitivity_header_row]:
    if cell.column <= len(sensitivity_table.columns) + 1:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

# Label first column as WACC
ws.cell(
    row=sensitivity_header_row,
    column=1,
    value="WACC (%)"
)

ws.cell(
    row=sensitivity_header_row,
    column=1
).font = Font(
    bold=True,
    color="FFFFFF"
)

ws.cell(
    row=sensitivity_header_row,
    column=1
).fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

ws.cell(
    row=sensitivity_header_row,
    column=1
).alignment = Alignment(
    horizontal="center",
    vertical="center"
)

# Sensitivity Analysis data
sensitivity_data_start = sensitivity_header_row + 1

sensitivity_data_end = (
    sensitivity_data_start
    + len(sensitivity_table)
    - 1
)

for row in ws.iter_rows(
    min_row=sensitivity_data_start,
    max_row=sensitivity_data_end,
    min_col=1,
    max_col=len(sensitivity_table.columns) + 1
):
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.number_format = "0.00"


# DCF Summary header

summary_header_row = len(df_D) + 10

for cell in ws[summary_header_row]:
    if cell.column <= 2:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# DCF Summary data
summary_data_start = summary_header_row + 1
summary_data_end = summary_data_start + len(summary_df) - 1

for row in ws.iter_rows(
    min_row=summary_data_start,
    max_row=summary_data_end,
    min_col=1,
    max_col=2
):
    for cell in row:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        if cell.column == 2:
            cell.number_format = "0.00"


# Bold DCF Summary section headings
ws.cell(
    row=summary_data_start,
    column=1
).font = Font(bold=True)

ws.cell(
    row=summary_data_start + 9,
    column=1
).font = Font(bold=True)


# Auto-adjust column widths
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(
                max_length,
                len(str(cell.value))
            )

    ws.column_dimensions[column_letter].width = max_length + 4


# Save workbook
wb.save(OUTPUT_FILE_PATH)
